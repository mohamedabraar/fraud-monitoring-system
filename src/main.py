import io
import secrets
from datetime import datetime

from flask import (
    Flask,
    render_template,
    request,
    send_file
)

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from alert_manager import create_alert
from database import (
    get_all_transactions,
    get_audit_events,
    get_dashboard_metrics,
    get_recent_alerts,
    get_recent_transactions,
    initialise_database,
    save_audit_event,
    save_transaction
)
from fraud_engine import calculate_risk
from transaction import Transaction
from validator import validate_transaction_data


app = Flask(__name__)

initialise_database()


def generate_transaction_id():
    return f"TX-{secrets.token_hex(3).upper()}"


@app.route("/", methods=["GET", "POST"])
def dashboard():
    assessment_result = None
    errors = []
    selected_audit_events = []

    if request.method == "POST":
        errors, amount = validate_transaction_data(request.form)

        if not errors:
            transaction = Transaction.create(
                transaction_id=generate_transaction_id(),
                customer_name=request.form["customer_name"],
                account_number=request.form["account_number"],
                branch=request.form["branch"],
                transaction_type=request.form["transaction_type"],
                channel=request.form["channel"],
                country=request.form["country"],
                amount=amount
            )

            assessment = calculate_risk(transaction)

            alert = create_alert(
                transaction,
                assessment
            )

            save_transaction(
                transaction,
                assessment,
                alert
            )

            save_audit_event(
                transaction.transaction_id,
                transaction.transaction_date,
                "Transaction submitted",
                "Transaction details were received and validated."
            )

            save_audit_event(
                transaction.transaction_id,
                transaction.transaction_date,
                "Risk assessment completed",
                (
                    f"Risk score {assessment.risk_score}/100 "
                    f"and risk level {assessment.risk_level} "
                    "were calculated."
                )
            )

            save_audit_event(
                transaction.transaction_id,
                alert.generated_at,
                "Alert decision completed",
                (
                    f"Alert status set to "
                    f"{alert.alert_status}."
                )
            )

            if alert.alert_status == "ACTIVE":
                save_audit_event(
                    transaction.transaction_id,
                    alert.generated_at,
                    "Fraud case created",
                    (
                        f"Case assigned to "
                        f"{alert.assigned_team}."
                    )
                )

            selected_audit_events = get_audit_events(
                transaction.transaction_id
            )

            assessment_result = {
                "transaction": transaction,
                "assessment": assessment,
                "alert": alert
            }

    metrics = get_dashboard_metrics()
    recent_transactions = get_recent_transactions(12)
    recent_alerts = get_recent_alerts(6)

    return render_template(
        "index.html",
        result=assessment_result,
        errors=errors,
        metrics=metrics,
        transactions=recent_transactions,
        alerts=recent_alerts,
        audit_events=selected_audit_events
    )


@app.route("/export/transactions.xlsx")
def export_transactions_excel():
    transactions = get_all_transactions()

    workbook = Workbook()

    report_sheet = workbook.active
    report_sheet.title = "Transactions"

    summary_sheet = workbook.create_sheet("Summary")

    # Corporate report colours
    navy = "051B35"
    blue = "0050B5"
    light_blue = "D9EAF7"
    white = "FFFFFF"
    dark_text = "172033"
    grey_text = "6D778B"
    border_colour = "D9E1EC"

    green = "C6EFCE"
    green_text = "006100"

    amber = "FFEB9C"
    amber_text = "9C6500"

    red = "FFC7CE"
    red_text = "9C0006"

    purple = "E4D3F5"
    purple_text = "7030A0"

    thin_border = Border(
        left=Side(style="thin", color=border_colour),
        right=Side(style="thin", color=border_colour),
        top=Side(style="thin", color=border_colour),
        bottom=Side(style="thin", color=border_colour)
    )

    # Report title
    report_sheet.merge_cells("A1:O1")

    title_cell = report_sheet["A1"]
    title_cell.value = "FRAUD MONITORING SYSTEM"
    title_cell.font = Font(
        name="Aptos Display",
        size=20,
        bold=True,
        color=white
    )
    title_cell.fill = PatternFill(
        fill_type="solid",
        fgColor=navy
    )
    title_cell.alignment = Alignment(
        horizontal="left",
        vertical="center"
    )

    report_sheet.row_dimensions[1].height = 34

    # Report subtitle
    report_sheet.merge_cells("A2:O2")

    subtitle_cell = report_sheet["A2"]
    subtitle_cell.value = "Transaction Assessment Report"
    subtitle_cell.font = Font(
        name="Aptos",
        size=12,
        bold=True,
        color=blue
    )
    subtitle_cell.fill = PatternFill(
        fill_type="solid",
        fgColor=light_blue
    )
    subtitle_cell.alignment = Alignment(
        horizontal="left",
        vertical="center"
    )

    report_sheet.row_dimensions[2].height = 24

    # Export metadata
    report_sheet["A3"] = "Generated:"
    report_sheet["B3"] = datetime.now().strftime(
        "%d %b %Y %H:%M:%S"
    )

    report_sheet["D3"] = "Total transactions:"
    report_sheet["E3"] = len(transactions)

    for cell_reference in ["A3", "D3"]:
        report_sheet[cell_reference].font = Font(
            bold=True,
            color=grey_text
        )

    # Column headers
    headers = [
        "Transaction ID",
        "Customer",
        "Account",
        "Branch",
        "Transaction Type",
        "Channel",
        "Country",
        "Amount",
        "Date",
        "Risk Level",
        "Risk Score",
        "Fraud Probability",
        "Alert Status",
        "Alert Reference",
        "Assigned Team"
    ]

    header_row = 5

    for column_number, header in enumerate(headers, start=1):
        cell = report_sheet.cell(
            row=header_row,
            column=column_number,
            value=header
        )

        cell.font = Font(
            name="Aptos",
            bold=True,
            color=white
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor=blue
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        cell.border = thin_border

    report_sheet.row_dimensions[header_row].height = 32

    # Transaction records
    first_data_row = header_row + 1

    for row_number, transaction in enumerate(
        transactions,
        start=first_data_row
    ):
        values = [
            transaction["transaction_id"],
            transaction["customer_name"],
            transaction["account_number"],
            transaction["branch"],
            transaction["transaction_type"],
            transaction["channel"],
            transaction["country"],
            transaction["amount"],
            transaction["transaction_date"],
            transaction["risk_level"],
            transaction["risk_score"],
            transaction["fraud_probability"],
            transaction["alert_status"],
            transaction["alert_reference"],
            transaction["assigned_team"]
        ]

        for column_number, value in enumerate(values, start=1):
            cell = report_sheet.cell(
                row=row_number,
                column=column_number,
                value=value
            )

            cell.font = Font(
                name="Aptos",
                size=10,
                color=dark_text
            )

            cell.alignment = Alignment(
                vertical="center"
            )

            cell.border = thin_border

            if row_number % 2 == 0:
                cell.fill = PatternFill(
                    fill_type="solid",
                    fgColor="F3F7FA"
                )

        # Amount format
        report_sheet.cell(
            row=row_number,
            column=8
        ).number_format = 'R #,##0.00'

        # Risk score format
        report_sheet.cell(
            row=row_number,
            column=11
        ).number_format = '0"/100"'

        # Fraud probability format
        report_sheet.cell(
            row=row_number,
            column=12
        ).number_format = '0"%"'

        # Risk-level colour
        risk_level = str(
            transaction["risk_level"]
        ).upper()

        risk_cell = report_sheet.cell(
            row=row_number,
            column=10
        )

        if risk_level == "LOW":
            risk_cell.fill = PatternFill(
                fill_type="solid",
                fgColor=green
            )
            risk_cell.font = Font(
                bold=True,
                color=green_text
            )

        elif risk_level == "MEDIUM":
            risk_cell.fill = PatternFill(
                fill_type="solid",
                fgColor=amber
            )
            risk_cell.font = Font(
                bold=True,
                color=amber_text
            )

        elif risk_level == "HIGH":
            risk_cell.fill = PatternFill(
                fill_type="solid",
                fgColor=red
            )
            risk_cell.font = Font(
                bold=True,
                color=red_text
            )

        elif risk_level == "CRITICAL":
            risk_cell.fill = PatternFill(
                fill_type="solid",
                fgColor=purple
            )
            risk_cell.font = Font(
                bold=True,
                color=purple_text
            )

        # Alert-status colour
        alert_cell = report_sheet.cell(
            row=row_number,
            column=13
        )

        alert_status = str(
            transaction["alert_status"]
        ).upper()

        if alert_status == "ACTIVE":
            alert_cell.fill = PatternFill(
                fill_type="solid",
                fgColor=red
            )
            alert_cell.font = Font(
                bold=True,
                color=red_text
            )

        elif alert_status == "MONITOR":
            alert_cell.fill = PatternFill(
                fill_type="solid",
                fgColor=amber
            )
            alert_cell.font = Font(
                bold=True,
                color=amber_text
            )

        elif alert_status == "CLEARED":
            alert_cell.fill = PatternFill(
                fill_type="solid",
                fgColor=green
            )
            alert_cell.font = Font(
                bold=True,
                color=green_text
            )

    # Create an Excel table when data exists
    if transactions:
        last_data_row = first_data_row + len(transactions) - 1

        transaction_table = Table(
            displayName="FraudTransactions",
            ref=f"A{header_row}:O{last_data_row}"
        )

        table_style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=False,
            showColumnStripes=False
        )

        transaction_table.tableStyleInfo = table_style
        report_sheet.add_table(transaction_table)

    # Column widths
    column_widths = {
        "A": 18,
        "B": 22,
        "C": 16,
        "D": 18,
        "E": 24,
        "F": 18,
        "G": 20,
        "H": 18,
        "I": 23,
        "J": 15,
        "K": 14,
        "L": 18,
        "M": 15,
        "N": 22,
        "O": 24
    }

    for column_letter, width in column_widths.items():
        report_sheet.column_dimensions[column_letter].width = width

    report_sheet.freeze_panes = "A6"
    report_sheet.auto_filter.ref = (
        f"A{header_row}:O"
        f"{header_row + max(len(transactions), 1)}"
    )

    report_sheet.sheet_view.showGridLines = False
    report_sheet.page_setup.orientation = "landscape"
    report_sheet.page_setup.fitToWidth = 1
    report_sheet.page_margins.left = 0.25
    report_sheet.page_margins.right = 0.25

    # Summary sheet
    summary_sheet.merge_cells("A1:D1")

    summary_title = summary_sheet["A1"]
    summary_title.value = "FRAUD MONITORING SUMMARY"
    summary_title.font = Font(
        name="Aptos Display",
        size=18,
        bold=True,
        color=white
    )
    summary_title.fill = PatternFill(
        fill_type="solid",
        fgColor=navy
    )
    summary_title.alignment = Alignment(
        horizontal="left",
        vertical="center"
    )

    summary_sheet.row_dimensions[1].height = 34

    total_value = sum(
        float(transaction["amount"])
        for transaction in transactions
    )

    active_alerts = sum(
        1
        for transaction in transactions
        if str(transaction["alert_status"]).upper() == "ACTIVE"
    )

    critical_transactions = sum(
        1
        for transaction in transactions
        if str(transaction["risk_level"]).upper() == "CRITICAL"
    )

    average_risk_score = (
        round(
            sum(
                int(transaction["risk_score"])
                for transaction in transactions
            ) / len(transactions)
        )
        if transactions
        else 0
    )

    summary_rows = [
        ("Report generated", datetime.now().strftime(
            "%d %b %Y %H:%M:%S"
        )),
        ("Total transactions", len(transactions)),
        ("Total value assessed", total_value),
        ("Active alerts", active_alerts),
        ("Critical transactions", critical_transactions),
        ("Average risk score", average_risk_score)
    ]

    for row_number, summary_row in enumerate(
        summary_rows,
        start=3
    ):
        label_cell = summary_sheet.cell(
            row=row_number,
            column=1,
            value=summary_row[0]
        )

        value_cell = summary_sheet.cell(
            row=row_number,
            column=2,
            value=summary_row[1]
        )

        label_cell.font = Font(
            bold=True,
            color=grey_text
        )

        value_cell.font = Font(
            bold=True,
            color=dark_text
        )

        label_cell.fill = PatternFill(
            fill_type="solid",
            fgColor="F3F7FA"
        )

        label_cell.border = thin_border
        value_cell.border = thin_border

        label_cell.alignment = Alignment(
            vertical="center"
        )

        value_cell.alignment = Alignment(
            vertical="center"
        )

    summary_sheet["B5"].number_format = 'R #,##0.00'
    summary_sheet["B8"].number_format = '0"/100"'

    summary_sheet.column_dimensions["A"].width = 28
    summary_sheet.column_dimensions["B"].width = 25
    summary_sheet.sheet_view.showGridLines = False

    # Save workbook to memory
    excel_output = io.BytesIO()
    workbook.save(excel_output)
    excel_output.seek(0)

    filename = (
        "fraud_transaction_report_"
        f"{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    )

    return send_file(
        excel_output,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )


if __name__ == "__main__":
    app.run(
        debug=True,
        host="127.0.0.1",
        port=5000
    )